"""Deterministic, inert DOCX export for source-backed research reports.

Only canonical Markdown produced by ``research.markdown_artifact`` is accepted.
The pinned ``python-docx`` dependency creates plain paragraphs, after which the
OOXML ZIP is stripped to a closed safe part set and repacked with fixed
metadata.  Inspection never invokes Office and rejects macros, embedded
objects, media, templates, hyperlinks, and external relationships.
"""

from __future__ import annotations

import hashlib
import io
import json
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET

from research.markdown_artifact import (
    ResearchMarkdownSchema,
    inspect_research_markdown,
)
from research.models import (
    MAX_RESEARCH_BYTES,
    MAX_RESEARCH_ROWS,
    ResearchValidationError,
    canonicalize_public_url,
)


DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)
MAX_RESEARCH_DOCX_BYTES = 8 * 1024 * 1024
MAX_RESEARCH_DOCX_TITLE_CHARS = 255
MAX_RESEARCH_DOCX_ENTRIES = 32
MAX_RESEARCH_DOCX_PART_BYTES = 4 * 1024 * 1024
MAX_RESEARCH_DOCX_UNCOMPRESSED_BYTES = 16 * 1024 * 1024
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FIXED_CORE_TIME = datetime(2000, 1, 1, 0, 0, 0)
_SHA256 = re.compile(r"^[0-9a-f]{64}$")
_SOURCE_TEXT = re.compile(
    r"^(S[0-9]{3}): "
    r"(https://[A-Za-z0-9._~:/?#[\]@!$&'()*+,;=%-]+)$"
)
_CITATION_TEXT = re.compile(
    r"^(S[0-9]{3}): "
    r"(https://[A-Za-z0-9._~:/?#[\]@!$&'()*+,;=%-]+)$"
)
_RELATIONSHIPS_NS = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
_CONTENT_TYPES_NS = (
    "http://schemas.openxmlformats.org/package/2006/content-types"
)
_WORD_NS = (
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
)
_RELATIONSHIP_TAG = f"{{{_RELATIONSHIPS_NS}}}Relationship"
_WORD_PARAGRAPH = f"{{{_WORD_NS}}}p"
_WORD_TEXT = f"{{{_WORD_NS}}}t"
_WORD_STYLE = f"{{{_WORD_NS}}}pStyle"
_WORD_STYLE_VALUE = f"{{{_WORD_NS}}}val"
_REQUIRED_PARTS = frozenset(
    {
        "[Content_Types].xml",
        "_rels/.rels",
        "word/document.xml",
    }
)
_ALLOWED_PARTS = frozenset(
    {
        *_REQUIRED_PARTS,
        "docProps/app.xml",
        "docProps/core.xml",
        "word/_rels/document.xml.rels",
        "word/fontTable.xml",
        "word/numbering.xml",
        "word/settings.xml",
        "word/styles.xml",
        "word/stylesWithEffects.xml",
        "word/theme/theme1.xml",
        "word/webSettings.xml",
    }
)
_REMOVED_PART_PREFIXES = (
    "customXml/",
    "docProps/thumbnail.",
)
_DANGEROUS_PART_MARKERS = (
    "activex",
    "altchunk",
    "attachments",
    "embeddings",
    "externallinks",
    "macrosheets",
    "oleobject",
    "vbaproject",
)
_DANGEROUS_RELATIONSHIP_SUFFIXES = (
    "/attachedtemplate",
    "/audiovideo",
    "/externallink",
    "/hyperlink",
    "/oleobject",
    "/package",
)
_ALLOWED_RELATIONSHIP_SUFFIXES = (
    "/core-properties",
    "/extended-properties",
    "/fonttable",
    "/numbering",
    "/officedocument",
    "/settings",
    "/styles",
    "/styleswitheffects",
    "/theme",
    "/websettings",
)
_DANGEROUS_XML_MARKERS = (
    b"<!doctype",
    b"<!entity",
    b"<w:altchunk",
    b"<w:object",
    b"<w:pict",
    b"<w:subdoc",
    b"<w:hyperlink",
)


class ResearchDocxError(ValueError):
    """A DOCX report request, package, or extracted document is invalid."""


@dataclass(frozen=True, slots=True)
class ResearchDocxParagraph:
    style: str
    text: str = field(repr=False)

    def __post_init__(self) -> None:
        if self.style not in {"Title", "Heading1", "Body"}:
            raise ResearchDocxError(
                "Research DOCX paragraph style is invalid"
            )
        if (
            not isinstance(self.text, str)
            or not self.text
            or len(self.text) > 16 * 1024
            or not self.text.isprintable()
        ):
            raise ResearchDocxError(
                "Research DOCX paragraph text is invalid"
            )


@dataclass(frozen=True, slots=True)
class ResearchDocxModel:
    paragraphs: tuple[ResearchDocxParagraph, ...]
    report_sha256: str
    document_digest: str
    source_digest: str
    section_count: int

    def __post_init__(self) -> None:
        if (
            not isinstance(self.paragraphs, tuple)
            or not self.paragraphs
            or any(
                not isinstance(item, ResearchDocxParagraph)
                for item in self.paragraphs
            )
        ):
            raise TypeError(
                "Research DOCX model paragraphs must be typed"
            )
        for digest in (
            self.report_sha256,
            self.document_digest,
            self.source_digest,
        ):
            if not isinstance(digest, str) or _SHA256.fullmatch(digest) is None:
                raise ResearchDocxError(
                    "Research DOCX model digest is invalid"
                )
        _bounded_integer(
            self.section_count,
            1,
            MAX_RESEARCH_ROWS,
            "Research DOCX model section count",
        )
        if self.document_digest != _paragraph_digest(self.paragraphs):
            raise ResearchDocxError(
                "Research DOCX model paragraph digest is invalid"
            )


@dataclass(frozen=True, slots=True)
class ResearchDocxArtifact:
    content: bytes = field(repr=False)
    sha256: str
    report_sha256: str
    document_digest: str
    source_digest: str
    section_count: int

    def __post_init__(self) -> None:
        if (
            not isinstance(self.content, bytes)
            or not 1 <= len(self.content) <= MAX_RESEARCH_DOCX_BYTES
            or self.sha256 != hashlib.sha256(self.content).hexdigest()
        ):
            raise ResearchDocxError(
                "Research DOCX artifact content is invalid"
            )
        for digest in (
            self.sha256,
            self.report_sha256,
            self.document_digest,
            self.source_digest,
        ):
            if not isinstance(digest, str) or _SHA256.fullmatch(digest) is None:
                raise ResearchDocxError(
                    "Research DOCX artifact digest is invalid"
                )
        _bounded_integer(
            self.section_count,
            1,
            MAX_RESEARCH_ROWS,
            "Research DOCX artifact section count",
        )


@dataclass(frozen=True, slots=True)
class ResearchDocxInspection:
    structurally_valid: bool
    section_count: int
    citation_count: int
    paragraph_count: int
    package_part_count: int
    content_sha256: str
    document_digest: str
    source_digest: str
    result_code: str

    def __post_init__(self) -> None:
        if type(self.structurally_valid) is not bool:
            raise TypeError(
                "Research DOCX inspection validity must be explicit"
            )
        for value, maximum, label in (
            (self.section_count, MAX_RESEARCH_ROWS, "section"),
            (self.citation_count, 10_000, "citation"),
            (self.paragraph_count, 10_000, "paragraph"),
            (
                self.package_part_count,
                MAX_RESEARCH_DOCX_ENTRIES,
                "package part",
            ),
        ):
            if type(value) is not int or not 0 <= value <= maximum:
                raise ResearchDocxError(
                    f"Research DOCX inspection {label} count is invalid"
                )
        for digest in (
            self.content_sha256,
            self.document_digest,
            self.source_digest,
        ):
            if not isinstance(digest, str) or _SHA256.fullmatch(digest) is None:
                raise ResearchDocxError(
                    "Research DOCX inspection digest is invalid"
                )
        if self.result_code not in {
            "invalid_research_docx",
            "research_docx_complete",
        }:
            raise ResearchDocxError(
                "Research DOCX inspection result is invalid"
            )
        if self.structurally_valid != (
            self.result_code == "research_docx_complete"
        ):
            raise ResearchDocxError(
                "Research DOCX inspection outcome is inconsistent"
            )


def research_markdown_to_docx_model(
    report_content: bytes,
    schema: ResearchMarkdownSchema,
    *,
    requested_sections: int,
    maximum_report_bytes: int = MAX_RESEARCH_BYTES,
) -> ResearchDocxModel:
    """Convert only a complete canonical report into styled plain text."""

    if len(schema.title) > MAX_RESEARCH_DOCX_TITLE_CHARS:
        raise ResearchDocxError(
            "Research DOCX title is too long"
        )
    inspection = inspect_research_markdown(
        report_content,
        schema,
        requested_sections=requested_sections,
        maximum_bytes=maximum_report_bytes,
    )
    if not inspection.complete:
        raise ResearchDocxError(
            "Research DOCX source report is not complete and canonical"
        )
    text = report_content.decode("utf-8")
    lines = text.splitlines()
    paragraphs: list[ResearchDocxParagraph] = []
    for line in lines:
        if not line:
            continue
        if line.startswith("# "):
            paragraphs.append(
                ResearchDocxParagraph(
                    "Title",
                    _unescape_markdown(line[2:]),
                )
            )
            continue
        if line.startswith("## "):
            paragraphs.append(
                ResearchDocxParagraph(
                    "Heading1",
                    _unescape_markdown(line[3:]),
                )
            )
            continue
        if line.startswith("- **"):
            label, remainder = line[4:].split(":** ", 1)
            if " — " in remainder:
                raw_value, raw_citations = remainder.rsplit(" — ", 1)
                citations = "; ".join(
                    _docx_citation_text(item)
                    for item in raw_citations.split("; ")
                )
                rendered = (
                    f"{label}: {_unescape_markdown(raw_value)}"
                    f" — {citations}"
                )
            else:
                rendered = (
                    f"{label}: "
                    f"{remainder}"
                )
            paragraphs.append(
                ResearchDocxParagraph("Body", rendered)
            )
            continue
        match = re.fullmatch(
            r"- \[(S[0-9]{3})\] <([^<>\s]+)>",
            line,
        )
        if match is None:
            raise ResearchDocxError(
                "Research DOCX source report line is invalid"
            )
        paragraphs.append(
            ResearchDocxParagraph(
                "Body",
                f"{match.group(1)}: {match.group(2)}",
            )
        )
    model = ResearchDocxModel(
        paragraphs=tuple(paragraphs),
        report_sha256=hashlib.sha256(report_content).hexdigest(),
        document_digest=_paragraph_digest(tuple(paragraphs)),
        source_digest=inspection.source_digest,
        section_count=inspection.section_count,
    )
    _validate_docx_paragraphs(
        model.paragraphs,
        schema,
        requested_sections=requested_sections,
    )
    return model


def render_research_markdown_to_docx(
    report_content: bytes,
    schema: ResearchMarkdownSchema,
    *,
    requested_sections: int,
    maximum_output_bytes: int,
) -> ResearchDocxArtifact:
    """Render one deterministic macro-free DOCX with pinned python-docx."""

    _bounded_integer(
        maximum_output_bytes,
        1,
        MAX_RESEARCH_DOCX_BYTES,
        "Research DOCX output limit",
    )
    model = research_markdown_to_docx_model(
        report_content,
        schema,
        requested_sections=requested_sections,
    )
    try:
        from docx import Document
    except ImportError as exc:
        raise ResearchDocxError(
            "Pinned python-docx runtime is unavailable"
        ) from exc

    document = Document()
    properties = document.core_properties
    properties.author = "Clicky"
    properties.comments = ""
    properties.created = _FIXED_CORE_TIME
    properties.keywords = ""
    properties.last_modified_by = "Clicky"
    properties.modified = _FIXED_CORE_TIME
    properties.revision = 1
    properties.subject = "Verified source-backed research report"
    properties.title = schema.title
    for paragraph in model.paragraphs:
        if paragraph.style == "Title":
            document.add_paragraph(paragraph.text, style="Title")
        elif paragraph.style == "Heading1":
            document.add_heading(paragraph.text, level=1)
        else:
            document.add_paragraph(paragraph.text)
    raw = io.BytesIO()
    document.save(raw)
    content = _canonicalize_docx_package(raw.getvalue())
    if len(content) > maximum_output_bytes:
        raise ResearchDocxError(
            "Research DOCX output limit exceeded"
        )
    inspection = inspect_research_docx(
        content,
        schema,
        requested_sections=requested_sections,
        maximum_bytes=maximum_output_bytes,
    )
    if (
        not inspection.structurally_valid
        or inspection.document_digest != model.document_digest
        or inspection.source_digest != model.source_digest
    ):
        raise ResearchDocxError(
            "Rendered Research DOCX failed exact self-verification"
        )
    return ResearchDocxArtifact(
        content=content,
        sha256=hashlib.sha256(content).hexdigest(),
        report_sha256=model.report_sha256,
        document_digest=model.document_digest,
        source_digest=model.source_digest,
        section_count=model.section_count,
    )


def validate_safe_docx_package(
    content: bytes,
    *,
    maximum_bytes: int = MAX_RESEARCH_DOCX_BYTES,
) -> tuple[str, ...]:
    """Validate a closed inert OOXML container without extracting report text."""

    entries = _read_safe_zip_entries(
        content,
        maximum_bytes=maximum_bytes,
    )
    _validate_package_parts(entries)
    _validate_content_types(entries["[Content_Types].xml"])
    _validate_relationships(entries)
    _parse_xml(entries["word/document.xml"], "word/document.xml")
    return tuple(sorted(entries))


def safe_docx_text_preview(
    content: bytes,
    *,
    maximum_chars: int = 24 * 1024,
) -> str:
    """Return bounded extracted plain text from an already inert package."""

    _bounded_integer(
        maximum_chars,
        1,
        64 * 1024,
        "Research DOCX preview character limit",
    )
    entries = _read_safe_zip_entries(
        content,
        maximum_bytes=MAX_RESEARCH_DOCX_BYTES,
    )
    _validate_package_parts(entries)
    _validate_content_types(entries["[Content_Types].xml"])
    _validate_relationships(entries)
    paragraphs = _extract_document_paragraphs(
        entries["word/document.xml"]
    )
    return "\n".join(item.text for item in paragraphs)[:maximum_chars]


def inspect_research_docx(
    content: bytes,
    schema: ResearchMarkdownSchema,
    *,
    requested_sections: int,
    maximum_bytes: int,
) -> ResearchDocxInspection:
    """Inspect safe OOXML and extracted source coverage without Office."""

    if not isinstance(content, bytes):
        raise TypeError(
            "Research DOCX inspection requires immutable bytes"
        )
    if not isinstance(schema, ResearchMarkdownSchema):
        raise TypeError(
            "Research DOCX inspection requires a typed report schema"
        )
    _bounded_integer(
        requested_sections,
        1,
        MAX_RESEARCH_ROWS,
        "Research DOCX requested section count",
    )
    _bounded_integer(
        maximum_bytes,
        1,
        MAX_RESEARCH_DOCX_BYTES,
        "Research DOCX verification byte limit",
    )
    content_digest = hashlib.sha256(content).hexdigest()
    empty_digest = hashlib.sha256(b"").hexdigest()
    section_count = 0
    citation_count = 0
    paragraph_count = 0
    part_count = 0
    document_digest = empty_digest
    source_digest = _digest_json({"sources": ()})
    try:
        entries = _read_safe_zip_entries(
            content,
            maximum_bytes=maximum_bytes,
        )
        part_count = len(entries)
        _validate_package_parts(entries)
        _validate_content_types(entries["[Content_Types].xml"])
        _validate_relationships(entries)
        paragraphs = _extract_document_paragraphs(
            entries["word/document.xml"]
        )
        paragraph_count = len(paragraphs)
        section_count, citation_count, sources = _validate_docx_paragraphs(
            paragraphs,
            schema,
            requested_sections=requested_sections,
        )
        document_digest = _paragraph_digest(paragraphs)
        source_digest = _digest_json({"sources": sources})
    except (
        ET.ParseError,
        OSError,
        ResearchDocxError,
        ResearchValidationError,
        ValueError,
        zipfile.BadZipFile,
    ):
        return ResearchDocxInspection(
            structurally_valid=False,
            section_count=min(section_count, MAX_RESEARCH_ROWS),
            citation_count=min(citation_count, 10_000),
            paragraph_count=min(paragraph_count, 10_000),
            package_part_count=min(
                part_count,
                MAX_RESEARCH_DOCX_ENTRIES,
            ),
            content_sha256=content_digest,
            document_digest=document_digest,
            source_digest=source_digest,
            result_code="invalid_research_docx",
        )
    return ResearchDocxInspection(
        structurally_valid=True,
        section_count=section_count,
        citation_count=citation_count,
        paragraph_count=paragraph_count,
        package_part_count=part_count,
        content_sha256=content_digest,
        document_digest=document_digest,
        source_digest=source_digest,
        result_code="research_docx_complete",
    )


def _canonicalize_docx_package(content: bytes) -> bytes:
    entries = _read_safe_zip_entries(
        content,
        maximum_bytes=MAX_RESEARCH_DOCX_BYTES,
        allow_unreviewed_parts=True,
    )
    retained = {
        name: value
        for name, value in entries.items()
        if not name.startswith(_REMOVED_PART_PREFIXES)
    }
    if "_rels/.rels" in retained:
        retained["_rels/.rels"] = _remove_relationship_targets(
            retained["_rels/.rels"],
            _REMOVED_PART_PREFIXES,
        )
    if "word/_rels/document.xml.rels" in retained:
        retained["word/_rels/document.xml.rels"] = (
            _remove_relationship_targets(
                retained["word/_rels/document.xml.rels"],
                ("../customXml/",),
            )
        )
    retained["[Content_Types].xml"] = _remove_content_type_parts(
        retained["[Content_Types].xml"],
        _REMOVED_PART_PREFIXES,
    )
    _validate_package_parts(retained)
    stream = io.BytesIO()
    with zipfile.ZipFile(
        stream,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
        allowZip64=False,
    ) as archive:
        for name in sorted(retained):
            info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o100600 << 16
            archive.writestr(
                info,
                retained[name],
                compress_type=zipfile.ZIP_DEFLATED,
                compresslevel=9,
            )
    canonical = stream.getvalue()
    validate_safe_docx_package(canonical)
    return canonical


def _read_safe_zip_entries(
    content: bytes,
    *,
    maximum_bytes: int,
    allow_unreviewed_parts: bool = False,
) -> dict[str, bytes]:
    if (
        not isinstance(content, bytes)
        or not 1 <= len(content) <= maximum_bytes
    ):
        raise ResearchDocxError(
            "Research DOCX package size is invalid"
        )
    entries: dict[str, bytes] = {}
    normalized_names: set[str] = set()
    total = 0
    with zipfile.ZipFile(io.BytesIO(content), mode="r") as archive:
        infos = archive.infolist()
        if not 1 <= len(infos) <= MAX_RESEARCH_DOCX_ENTRIES:
            raise ResearchDocxError(
                "Research DOCX package part count is invalid"
            )
        for info in infos:
            name = info.filename
            _validate_part_name(name)
            normalized_name = name.casefold()
            mode = (info.external_attr >> 16) & 0o170000
            if (
                name in entries
                or normalized_name in normalized_names
                or info.is_dir()
                or info.flag_bits & 0x1
                or info.compress_type
                not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
                or mode not in {0, 0o100000}
            ):
                raise ResearchDocxError(
                    "Research DOCX package entry is invalid"
                )
            if (
                info.file_size > MAX_RESEARCH_DOCX_PART_BYTES
                or info.compress_size > MAX_RESEARCH_DOCX_PART_BYTES
                or (
                    info.compress_size == 0
                    and info.file_size > 0
                )
                or (
                    info.compress_size
                    and info.file_size > info.compress_size * 100
                )
            ):
                raise ResearchDocxError(
                    "Research DOCX package expansion is invalid"
                )
            total += info.file_size
            if total > MAX_RESEARCH_DOCX_UNCOMPRESSED_BYTES:
                raise ResearchDocxError(
                    "Research DOCX uncompressed size is invalid"
                )
            entries[name] = archive.read(info)
            normalized_names.add(normalized_name)
    if not allow_unreviewed_parts:
        _validate_package_parts(entries)
    return entries


def _validate_part_name(name: object) -> str:
    if (
        not isinstance(name, str)
        or not name
        or len(name) > 255
        or "\\" in name
        or name.startswith("/")
        or ":" in name
        or any(
            part in {"", ".", ".."}
            for part in PurePosixPath(name).parts
        )
    ):
        raise ResearchDocxError(
            "Research DOCX package path is invalid"
        )
    return name


def _validate_package_parts(entries: dict[str, bytes]) -> None:
    names = set(entries)
    if not _REQUIRED_PARTS.issubset(names):
        raise ResearchDocxError(
            "Research DOCX required package parts are missing"
        )
    if not names.issubset(_ALLOWED_PARTS):
        raise ResearchDocxError(
            "Research DOCX contains an unreviewed package part"
        )
    for name, content in entries.items():
        lowered = name.lower()
        if any(marker in lowered for marker in _DANGEROUS_PART_MARKERS):
            raise ResearchDocxError(
                "Research DOCX contains an active package part"
            )
        if name.endswith((".xml", ".rels")):
            _parse_xml(content, name)


def _parse_xml(content: bytes, name: str) -> ET.Element:
    if (
        not content
        or len(content) > MAX_RESEARCH_DOCX_PART_BYTES
        or any(marker in content.lower() for marker in _DANGEROUS_XML_MARKERS)
    ):
        raise ResearchDocxError(
            f"Research DOCX XML part is unsafe: {name}"
        )
    root = ET.fromstring(content)
    dangerous_local_names = {
        "altChunk",
        "attachedTemplate",
        "control",
        "externalLink",
        "hyperlink",
        "object",
        "oleObject",
        "pict",
        "subDoc",
    }
    if any(
        node.tag.rsplit("}", 1)[-1] in dangerous_local_names
        for node in root.iter()
    ):
        raise ResearchDocxError(
            f"Research DOCX XML part contains active content: {name}"
        )
    return root


def _validate_content_types(content: bytes) -> None:
    root = _parse_xml(content, "[Content_Types].xml")
    for node in root:
        content_type = node.attrib.get("ContentType", "").lower()
        part_name = node.attrib.get("PartName")
        if any(
            marker in content_type
            for marker in (
                "activex",
                "external",
                "macroenabled",
                "oleobject",
                "vba",
            )
        ):
            raise ResearchDocxError(
                "Research DOCX content type is active"
            )
        if part_name is not None:
            normalized = part_name.removeprefix("/")
            if normalized not in _ALLOWED_PARTS:
                raise ResearchDocxError(
                    "Research DOCX content type names an unknown part"
                )


def _validate_relationships(entries: dict[str, bytes]) -> None:
    for name, content in entries.items():
        if not name.endswith(".rels"):
            continue
        root = _parse_xml(content, name)
        relationship_ids: set[str] = set()
        for relationship in root.findall(_RELATIONSHIP_TAG):
            target = relationship.attrib.get("Target", "")
            relation_type = relationship.attrib.get("Type", "").lower()
            relationship_id = relationship.attrib.get("Id", "")
            if (
                not target
                or not relationship_id
                or relationship_id in relationship_ids
                or relationship.attrib.get("TargetMode") is not None
                or any(
                    relation_type.endswith(suffix)
                    for suffix in _DANGEROUS_RELATIONSHIP_SUFFIXES
                )
                or not any(
                    relation_type.endswith(suffix)
                    for suffix in _ALLOWED_RELATIONSHIP_SUFFIXES
                )
            ):
                raise ResearchDocxError(
                    "Research DOCX external or active relationship"
                )
            relationship_ids.add(relationship_id)
            resolved = _resolve_relationship_target(name, target)
            if resolved not in entries:
                raise ResearchDocxError(
                    "Research DOCX relationship target is missing"
                )


def _resolve_relationship_target(
    relationship_part: str,
    target: str,
) -> str:
    if (
        not target
        or "\\" in target
        or target.startswith("/")
        or ":" in target
    ):
        raise ResearchDocxError(
            "Research DOCX relationship target is invalid"
        )
    if relationship_part == "_rels/.rels":
        base = ""
    else:
        rels_parent = posixpath.dirname(relationship_part)
        source_parent = rels_parent.removesuffix("/_rels")
        base = source_parent
    resolved = posixpath.normpath(posixpath.join(base, target))
    if resolved.startswith("../") or resolved in {"", ".", ".."}:
        raise ResearchDocxError(
            "Research DOCX relationship escapes the package"
        )
    return resolved


def _extract_document_paragraphs(
    content: bytes,
) -> tuple[ResearchDocxParagraph, ...]:
    root = _parse_xml(content, "word/document.xml")
    forbidden_local_names = {
        "altChunk",
        "br",
        "drawing",
        "fldChar",
        "hyperlink",
        "instrText",
        "object",
        "pict",
        "subDoc",
        "tab",
    }
    for node in root.iter():
        local_name = node.tag.rsplit("}", 1)[-1]
        if local_name in forbidden_local_names:
            raise ResearchDocxError(
                "Research DOCX document contains active or hidden content"
            )
    paragraphs: list[ResearchDocxParagraph] = []
    body = root.find(f"{{{_WORD_NS}}}body")
    if body is None:
        raise ResearchDocxError(
            "Research DOCX document body is missing"
        )
    for node in body:
        if node.tag != _WORD_PARAGRAPH:
            continue
        style_node = node.find(
            f"{{{_WORD_NS}}}pPr/{_WORD_STYLE}"
        )
        raw_style = (
            style_node.attrib.get(_WORD_STYLE_VALUE)
            if style_node is not None
            else None
        )
        style = {
            None: "Body",
            "Normal": "Body",
            "Title": "Title",
            "Heading1": "Heading1",
        }.get(raw_style)
        if style is None:
            raise ResearchDocxError(
                "Research DOCX paragraph style is not permitted"
            )
        text = "".join(
            item.text or ""
            for item in node.iter(_WORD_TEXT)
        )
        paragraphs.append(ResearchDocxParagraph(style, text))
    if not paragraphs:
        raise ResearchDocxError(
            "Research DOCX document has no paragraphs"
        )
    return tuple(paragraphs)


def _validate_docx_paragraphs(
    paragraphs: tuple[ResearchDocxParagraph, ...],
    schema: ResearchMarkdownSchema,
    *,
    requested_sections: int,
) -> tuple[int, int, tuple[str, ...]]:
    if (
        not paragraphs
        or paragraphs[0]
        != ResearchDocxParagraph("Title", schema.title)
    ):
        raise ResearchDocxError(
            "Research DOCX title is invalid"
        )
    index = 1
    citation_count = 0
    cited_labels: set[str] = set()
    expected_labels = (
        "Entity",
        "Public URL",
        *schema.requested_field_ids,
        "Rationale",
    )
    for section in range(1, requested_sections + 1):
        if (
            index >= len(paragraphs)
            or paragraphs[index].style != "Heading1"
            or not paragraphs[index].text.startswith(f"{section}. ")
        ):
            raise ResearchDocxError(
                "Research DOCX section heading is invalid"
            )
        index += 1
        for label in expected_labels:
            if (
                index >= len(paragraphs)
                or paragraphs[index].style != "Body"
                or not paragraphs[index].text.startswith(f"{label}: ")
                or " — " not in paragraphs[index].text
            ):
                raise ResearchDocxError(
                    "Research DOCX factual paragraph is invalid"
                )
            raw_citations = paragraphs[index].text.rsplit(" — ", 1)[1]
            citations = raw_citations.split("; ")
            parsed = []
            for item in citations:
                match = _CITATION_TEXT.fullmatch(item)
                if match is None:
                    raise ResearchDocxError(
                        "Research DOCX citation is invalid"
                    )
                parsed.append(match.groups())
                cited_labels.add(match.group(1))
            if not parsed or parsed != sorted(set(parsed)):
                raise ResearchDocxError(
                    "Research DOCX citations are not canonical"
                )
            citation_count += len(parsed)
            index += 1
        if (
            index >= len(paragraphs)
            or paragraphs[index].style != "Body"
            or not paragraphs[index].text.startswith("Retrieved at: ")
        ):
            raise ResearchDocxError(
                "Research DOCX retrieval time is invalid"
            )
        timestamp = datetime.fromisoformat(
            paragraphs[index].text.removeprefix("Retrieved at: ")
        )
        if timestamp.tzinfo is None or timestamp.utcoffset() is None:
            raise ResearchDocxError(
                "Research DOCX retrieval time is invalid"
            )
        index += 1
    if (
        index >= len(paragraphs)
        or paragraphs[index]
        != ResearchDocxParagraph("Heading1", "Sources")
    ):
        raise ResearchDocxError(
            "Research DOCX sources heading is invalid"
        )
    index += 1
    sources: list[str] = []
    source_labels: set[str] = set()
    while index < len(paragraphs):
        paragraph = paragraphs[index]
        match = (
            _SOURCE_TEXT.fullmatch(paragraph.text)
            if paragraph.style == "Body"
            else None
        )
        if match is None:
            raise ResearchDocxError(
                "Research DOCX source paragraph is invalid"
            )
        label, raw_url = match.groups()
        if label != f"S{len(sources) + 1:03d}":
            raise ResearchDocxError(
                "Research DOCX source labels are not contiguous"
            )
        canonical = canonicalize_public_url(
            raw_url,
            require_https=True,
        )
        if canonical != raw_url or raw_url in sources:
            raise ResearchDocxError(
                "Research DOCX source URL is invalid"
            )
        sources.append(raw_url)
        source_labels.add(label)
        index += 1
    if not sources or source_labels != cited_labels:
        raise ResearchDocxError(
            "Research DOCX source coverage is incomplete"
        )
    return requested_sections, citation_count, tuple(sources)


def _docx_citation_text(value: str) -> str:
    match = re.fullmatch(
        r"\[(S[0-9]{3})\] <([^<>\s]+)>",
        value,
    )
    if match is None:
        raise ResearchDocxError(
            "Research DOCX source citation is invalid"
        )
    return f"{match.group(1)}: {match.group(2)}"


def _unescape_markdown(value: str) -> str:
    return re.sub(r"\\([!\"#$%&'()*+,\-./:;<=>?@[\\\]^_`{|}~])", r"\1", value)


def _paragraph_digest(
    paragraphs: tuple[ResearchDocxParagraph, ...],
) -> str:
    return _digest_json(
        {
            "paragraphs": [
                {"style": item.style, "text": item.text}
                for item in paragraphs
            ]
        }
    )


def _remove_relationship_targets(
    content: bytes,
    prefixes: tuple[str, ...],
) -> bytes:
    root = _parse_xml(content, "relationships")
    for node in tuple(root.findall(_RELATIONSHIP_TAG)):
        target = node.attrib.get("Target", "")
        normalized = target.removeprefix("/")
        if any(normalized.startswith(prefix) for prefix in prefixes):
            root.remove(node)
    ET.register_namespace("", _RELATIONSHIPS_NS)
    return ET.tostring(
        root,
        encoding="utf-8",
        xml_declaration=True,
    )


def _remove_content_type_parts(
    content: bytes,
    prefixes: tuple[str, ...],
) -> bytes:
    root = _parse_xml(content, "[Content_Types].xml")
    for node in tuple(root):
        part = node.attrib.get("PartName", "").removeprefix("/")
        extension = node.attrib.get("Extension", "").lower()
        if (
            any(part.startswith(prefix) for prefix in prefixes)
            or extension in {"jpeg", "jpg", "png", "gif", "bmp"}
        ):
            root.remove(node)
    ET.register_namespace("", _CONTENT_TYPES_NS)
    return ET.tostring(
        root,
        encoding="utf-8",
        xml_declaration=True,
    )


def _digest_json(value: object) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        ).encode("utf-8")
    ).hexdigest()


def validate_research_report_paragraphs(
    paragraphs: tuple[ResearchDocxParagraph, ...],
    schema: ResearchMarkdownSchema,
    *,
    requested_sections: int,
) -> tuple[int, int, tuple[str, ...]]:
    """Validate the shared visible report paragraph contract."""

    return _validate_docx_paragraphs(
        paragraphs,
        schema,
        requested_sections=requested_sections,
    )


def research_report_paragraph_digest(
    paragraphs: tuple[ResearchDocxParagraph, ...],
) -> str:
    """Return the canonical digest shared by DOCX and PDF renderers."""

    return _paragraph_digest(paragraphs)


def _bounded_integer(
    value: object,
    minimum: int,
    maximum: int,
    label: str,
) -> None:
    if type(value) is not int or not minimum <= value <= maximum:
        raise ResearchDocxError(f"{label} is invalid")


__all__ = [
    "DOCX_MEDIA_TYPE",
    "MAX_RESEARCH_DOCX_BYTES",
    "MAX_RESEARCH_DOCX_TITLE_CHARS",
    "ResearchDocxArtifact",
    "ResearchDocxError",
    "ResearchDocxInspection",
    "ResearchDocxModel",
    "ResearchDocxParagraph",
    "inspect_research_docx",
    "render_research_markdown_to_docx",
    "research_report_paragraph_digest",
    "research_markdown_to_docx_model",
    "safe_docx_text_preview",
    "validate_research_report_paragraphs",
    "validate_safe_docx_package",
]
