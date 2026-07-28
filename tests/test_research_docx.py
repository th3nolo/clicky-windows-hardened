"""Safe deterministic DOCX rendering and package inspection tests."""

from __future__ import annotations

import hashlib
import io
import unittest
import zipfile
from xml.etree import ElementTree as ET

from research.docx_artifact import (
    ResearchDocxError,
    ResearchDocxParagraph,
    inspect_research_docx,
    render_research_markdown_to_docx,
    research_markdown_to_docx_model,
    validate_safe_docx_package,
)
from research.markdown_artifact import (
    ResearchMarkdownSchema,
    render_research_markdown,
)
from tests.test_research_csv import batch, record


REL_NS = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
CT_NS = (
    "http://schemas.openxmlformats.org/package/2006/content-types"
)
W_NS = (
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
)


def _zip(entries: dict[str, bytes]) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(
        stream,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return stream.getvalue()


def _minimal_docx(
    paragraphs: tuple[ResearchDocxParagraph, ...],
    *,
    relationship_target_mode: str | None = None,
    extra_entries: dict[str, bytes] | None = None,
) -> bytes:
    ET.register_namespace("", CT_NS)
    content_types = ET.Element(f"{{{CT_NS}}}Types")
    ET.SubElement(
        content_types,
        f"{{{CT_NS}}}Default",
        {
            "Extension": "rels",
            "ContentType": (
                "application/vnd.openxmlformats-package."
                "relationships+xml"
            ),
        },
    )
    ET.SubElement(
        content_types,
        f"{{{CT_NS}}}Default",
        {"Extension": "xml", "ContentType": "application/xml"},
    )
    ET.SubElement(
        content_types,
        f"{{{CT_NS}}}Override",
        {
            "PartName": "/word/document.xml",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document.main+xml"
            ),
        },
    )

    ET.register_namespace("", REL_NS)
    relationships = ET.Element(f"{{{REL_NS}}}Relationships")
    attributes = {
        "Id": "rId1",
        "Type": (
            "http://schemas.openxmlformats.org/officeDocument/"
            "2006/relationships/officeDocument"
        ),
        "Target": "word/document.xml",
    }
    if relationship_target_mode is not None:
        attributes["TargetMode"] = relationship_target_mode
    ET.SubElement(
        relationships,
        f"{{{REL_NS}}}Relationship",
        attributes,
    )

    ET.register_namespace("w", W_NS)
    document = ET.Element(f"{{{W_NS}}}document")
    body = ET.SubElement(document, f"{{{W_NS}}}body")
    for item in paragraphs:
        paragraph = ET.SubElement(body, f"{{{W_NS}}}p")
        if item.style != "Body":
            properties = ET.SubElement(
                paragraph,
                f"{{{W_NS}}}pPr",
            )
            style = ET.SubElement(
                properties,
                f"{{{W_NS}}}pStyle",
            )
            style.set(f"{{{W_NS}}}val", item.style)
        run = ET.SubElement(paragraph, f"{{{W_NS}}}r")
        text = ET.SubElement(run, f"{{{W_NS}}}t")
        text.text = item.text
    ET.SubElement(body, f"{{{W_NS}}}sectPr")

    entries = {
        "[Content_Types].xml": ET.tostring(
            content_types,
            encoding="utf-8",
            xml_declaration=True,
        ),
        "_rels/.rels": ET.tostring(
            relationships,
            encoding="utf-8",
            xml_declaration=True,
        ),
        "word/document.xml": ET.tostring(
            document,
            encoding="utf-8",
            xml_declaration=True,
        ),
    }
    entries.update(extra_entries or {})
    return _zip(entries)


def _rewrite_part(content: bytes, name: str, mutate) -> bytes:
    with zipfile.ZipFile(io.BytesIO(content), "r") as archive:
        entries = {
            info.filename: archive.read(info)
            for info in archive.infolist()
        }
    entries[name] = mutate(entries[name])
    return _zip(entries)


class ResearchDocxArtifactTests(unittest.TestCase):
    def setUp(self) -> None:
        self.schema = ResearchMarkdownSchema(
            "Creator research",
            ("audience",),
        )
        self.report = render_research_markdown(
            batch(record("one")),
            self.schema,
            requested_sections=1,
        )
        self.model = research_markdown_to_docx_model(
            self.report.content,
            self.schema,
            requested_sections=1,
        )

    def test_canonical_markdown_maps_to_exact_source_complete_paragraphs(self):
        self.assertEqual(self.model.section_count, 1)
        self.assertEqual(
            self.model.report_sha256,
            self.report.sha256,
        )
        self.assertEqual(
            self.model.source_digest,
            self.report.source_digest,
        )
        self.assertEqual(
            self.model.paragraphs[0],
            ResearchDocxParagraph("Title", "Creator research"),
        )
        self.assertEqual(
            self.model.paragraphs[1],
            ResearchDocxParagraph("Heading1", "1. Creator one"),
        )
        self.assertEqual(
            self.model.paragraphs[-3],
            ResearchDocxParagraph("Heading1", "Sources"),
        )

    def test_minimal_safe_package_is_inspected_without_office(self):
        content = _minimal_docx(self.model.paragraphs)
        parts = validate_safe_docx_package(content)
        self.assertEqual(
            set(parts),
            {
                "[Content_Types].xml",
                "_rels/.rels",
                "word/document.xml",
            },
        )
        inspection = inspect_research_docx(
            content,
            self.schema,
            requested_sections=1,
            maximum_bytes=1024 * 1024,
        )
        self.assertTrue(inspection.structurally_valid)
        self.assertEqual(inspection.section_count, 1)
        self.assertEqual(
            inspection.document_digest,
            self.model.document_digest,
        )
        self.assertEqual(
            inspection.source_digest,
            self.model.source_digest,
        )
        self.assertEqual(
            inspection.content_sha256,
            hashlib.sha256(content).hexdigest(),
        )

    def test_package_rejects_external_relationships_macros_and_tampering(self):
        external = _minimal_docx(
            self.model.paragraphs,
            relationship_target_mode="External",
        )
        macro = _minimal_docx(
            self.model.paragraphs,
            extra_entries={"word/vbaProject.bin": b"macro"},
        )
        altered = list(self.model.paragraphs)
        altered[2] = ResearchDocxParagraph(
            "Body",
            "Entity: Changed — S001: https://creator.example/one",
        )
        changed = _minimal_docx(tuple(altered))
        for content in (external, macro, changed):
            with self.subTest(content=content[:24]):
                inspection = inspect_research_docx(
                    content,
                    self.schema,
                    requested_sections=1,
                    maximum_bytes=1024 * 1024,
                )
                if content is changed:
                    self.assertTrue(inspection.structurally_valid)
                    self.assertNotEqual(
                        inspection.document_digest,
                        self.model.document_digest,
                    )
                else:
                    self.assertFalse(inspection.structurally_valid)

    def test_incomplete_report_and_active_document_elements_are_rejected(self):
        short = render_research_markdown(
            batch(record("one")),
            self.schema,
            requested_sections=2,
        )
        with self.assertRaisesRegex(
            ResearchDocxError,
            "not complete",
        ):
            research_markdown_to_docx_model(
                short.content,
                self.schema,
                requested_sections=2,
            )

        active = _rewrite_part(
            _minimal_docx(self.model.paragraphs),
            "word/document.xml",
            lambda value: value.replace(
                b"</w:body>",
                b"<w:altChunk w:id=\"rId9\"/></w:body>",
            ),
        )
        self.assertFalse(
            inspect_research_docx(
                active,
                self.schema,
                requested_sections=1,
                maximum_bytes=1024 * 1024,
            ).structurally_valid
        )
        long_schema = ResearchMarkdownSchema(
            "x" * 256,
            ("audience",),
        )
        long_report = render_research_markdown(
            batch(record("one")),
            long_schema,
            requested_sections=1,
        )
        with self.assertRaisesRegex(
            ResearchDocxError,
            "title is too long",
        ):
            research_markdown_to_docx_model(
                long_report.content,
                long_schema,
                requested_sections=1,
            )

    def test_pinned_python_docx_renderer_is_deterministic_when_available(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx is Windows-runtime-only in this lane")
        first = render_research_markdown_to_docx(
            self.report.content,
            self.schema,
            requested_sections=1,
            maximum_output_bytes=1024 * 1024,
        )
        second = render_research_markdown_to_docx(
            self.report.content,
            self.schema,
            requested_sections=1,
            maximum_output_bytes=1024 * 1024,
        )
        self.assertEqual(first.content, second.content)
        self.assertEqual(first.report_sha256, self.report.sha256)
        self.assertEqual(
            first.document_digest,
            self.model.document_digest,
        )
        self.assertNotIn(b"vbaProject", first.content)
        self.assertNotIn(b"customXml", first.content)


if __name__ == "__main__":
    unittest.main()
