"""Deterministic formula-safe XLSX rendering and package inspection tests."""

from __future__ import annotations

import hashlib
import io
import unittest
import zipfile

from research.csv_artifact import ResearchCsvSchema, render_research_csv
from research.xlsx_artifact import (
    MAX_RESEARCH_XLSX_ENTRIES,
    ResearchXlsxError,
    inspect_research_xlsx,
    render_research_csv_to_xlsx,
    safe_xlsx_table_preview,
    validate_safe_xlsx_package,
)
from tests.test_research_csv import batch, record, rewrite


class ResearchXlsxArtifactTests(unittest.TestCase):
    def setUp(self) -> None:
        self.schema = ResearchCsvSchema(("audience",))
        self.csv = render_research_csv(
            batch(record("one")),
            self.schema,
            requested_rows=1,
        )
        self.artifact = render_research_csv_to_xlsx(
            self.csv.content,
            self.schema,
            requested_rows=1,
            maximum_output_bytes=1024 * 1024,
        )

    def test_renderer_is_deterministic_and_exactly_self_verified(self):
        second = render_research_csv_to_xlsx(
            self.csv.content,
            self.schema,
            requested_rows=1,
            maximum_output_bytes=1024 * 1024,
        )

        self.assertEqual(self.artifact.content, second.content)
        self.assertEqual(
            self.artifact.sha256,
            hashlib.sha256(self.artifact.content).hexdigest(),
        )
        self.assertEqual(
            self.artifact.source_csv_sha256,
            self.csv.sha256,
        )
        self.assertEqual(self.artifact.schema_digest, self.schema.schema_digest)
        self.assertEqual(self.artifact.row_count, 1)
        self.assertEqual(
            self.artifact.column_count,
            len(self.schema.columns),
        )

        with zipfile.ZipFile(io.BytesIO(self.artifact.content)) as package:
            self.assertEqual(
                len(package.infolist()),
                MAX_RESEARCH_XLSX_ENTRIES,
            )
            names = tuple(info.filename for info in package.infolist())
            self.assertNotIn("xl/sharedStrings.xml", names)
            self.assertNotIn("xl/externalLinks/externalLink1.xml", names)
            worksheet = package.read("xl/worksheets/sheet1.xml")
            self.assertNotIn(b"<f", worksheet)
            self.assertNotIn(b"<hyperlinks", worksheet)
            self.assertNotIn(b"hidden", worksheet)

    def test_inspection_and_preview_expose_only_bounded_verified_content(self):
        inspection = inspect_research_xlsx(
            self.artifact.content,
            self.schema,
            requested_rows=1,
            maximum_bytes=1024 * 1024,
        )
        self.assertTrue(inspection.structurally_valid)
        self.assertEqual(inspection.result_code, "research_xlsx_complete")
        self.assertEqual(inspection.row_count, 1)
        self.assertEqual(
            inspection.source_csv_sha256,
            self.csv.sha256,
        )
        self.assertEqual(
            inspection.table_digest,
            self.artifact.table_digest,
        )

        preview = safe_xlsx_table_preview(self.artifact.content)
        self.assertIn('"row_count":1', preview)
        self.assertIn('"entity"', preview)
        self.assertIn("Creator one", preview)
        validate_safe_xlsx_package(self.artifact.content)

    def test_formula_prefixes_and_incomplete_source_csv_are_rejected(self):
        formula_csv = rewrite(
            self.csv.content,
            lambda rows: rows[1].__setitem__(0, "=WEBSERVICE(A1)"),
        )
        with self.assertRaisesRegex(ResearchXlsxError, "not complete"):
            render_research_csv_to_xlsx(
                formula_csv,
                self.schema,
                requested_rows=1,
            )

        incomplete = render_research_csv(
            batch(record("one")),
            self.schema,
            requested_rows=2,
        )
        with self.assertRaisesRegex(ResearchXlsxError, "not complete"):
            render_research_csv_to_xlsx(
                incomplete.content,
                self.schema,
                requested_rows=2,
            )

    def test_formula_elements_external_parts_and_content_tampering_fail_closed(
        self,
    ):
        formula = _rewrite_package(
            self.artifact.content,
            mutate_sheet=lambda content: content.replace(
                b"<is><t>Creator one</t></is>",
                b"<f>1+1</f><v>2</v>",
                1,
            ),
        )
        external = _rewrite_package(
            self.artifact.content,
            extra_parts={
                "xl/externalLinks/externalLink1.xml": b"<externalLink/>"
            },
        )
        macro = _rewrite_package(
            self.artifact.content,
            extra_parts={"xl/vbaProject.bin": b"macro"},
        )
        escaped = _rewrite_package(
            self.artifact.content,
            extra_parts={"../escape.xml": b"escaped"},
        )
        hidden = _rewrite_package(
            self.artifact.content,
            mutate_workbook=lambda content: content.replace(
                b'state="visible"',
                b'state="hidden"',
                1,
            ),
        )
        changed = self.artifact.content.replace(
            b"Creator one",
            b"Creator two",
            1,
        )

        for content in (formula, external, macro, escaped, hidden, changed):
            with self.subTest(content_sha256=hashlib.sha256(content).hexdigest()):
                with self.assertRaises(ResearchXlsxError):
                    validate_safe_xlsx_package(content)
                inspection = inspect_research_xlsx(
                    content,
                    self.schema,
                    requested_rows=1,
                    maximum_bytes=1024 * 1024,
                )
                self.assertFalse(inspection.structurally_valid)
                self.assertEqual(
                    inspection.result_code,
                    "invalid_research_xlsx",
                )

    def test_openpyxl_reads_representative_cells_when_available(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is Windows-runtime-only in this lane")

        workbook = load_workbook(
            io.BytesIO(self.artifact.content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        try:
            self.assertEqual(workbook.sheetnames, ["Research"])
            worksheet = workbook["Research"]
            self.assertEqual(worksheet["A1"].value, "entity")
            self.assertEqual(worksheet["A2"].value, "Creator one")
            self.assertEqual(worksheet["E1"].value, "audience")
            self.assertEqual(
                worksheet["E2"].value,
                "Software developers",
            )
        finally:
            workbook.close()


def _rewrite_package(
    content: bytes,
    *,
    mutate_sheet=lambda value: value,
    mutate_workbook=lambda value: value,
    extra_parts: dict[str, bytes] | None = None,
) -> bytes:
    with zipfile.ZipFile(io.BytesIO(content), "r") as source:
        parts = [
            (
                info.filename,
                (
                    mutate_sheet(source.read(info))
                    if info.filename == "xl/worksheets/sheet1.xml"
                    else (
                        mutate_workbook(source.read(info))
                        if info.filename == "xl/workbook.xml"
                        else source.read(info)
                    )
                ),
            )
            for info in source.infolist()
        ]
    parts.extend((extra_parts or {}).items())
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as target:
        for name, value in parts:
            target.writestr(name, value)
    return output.getvalue()


if __name__ == "__main__":
    unittest.main()
